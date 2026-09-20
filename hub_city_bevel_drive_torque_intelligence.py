from io import BytesIO
from pathlib import Path
import re

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# =========================================================
# FILE SETTINGS
# =========================================================

APP_DIR = Path(__file__).resolve().parent

DATABASE_FILE = (
    APP_DIR / "HubCity_Bevel_Gear_Torque_Data.xlsx"
)

IN_LB_TO_NM = 0.112984829


# =========================================================
# WEBPAGE SETTINGS
# =========================================================

st.set_page_config(
    page_title="Hub City Bevel Drive Torque Intelligence",
    page_icon="⚙️",
    layout="wide",
)

st.title("⚙️ Hub City Bevel Drive Torque Intelligence")

st.caption(
    "Torque lookup, RPM rating table, and Excel report generation."
)


# =========================================================
# SUPPORTING FUNCTIONS
# =========================================================

def normalized_name(value):
    """
    Remove spaces, underscores and special characters.

    Example:
    Gear_Type -> GEARTYPE
    Gear Type -> GEARTYPE
    """

    return re.sub(
        r"[^A-Z0-9]",
        "",
        str(value).upper(),
    )


def find_column(dataframe, possible_names):
    """
    Find an Excel column even if its heading uses
    a slightly different format.
    """

    available_columns = {
        normalized_name(column): column
        for column in dataframe.columns
    }

    for name in possible_names:

        normalized = normalized_name(name)

        if normalized in available_columns:
            return available_columns[normalized]

    return None


def find_sheet(sheet_names, possible_names):
    """
    Find the required Excel worksheet even if the
    worksheet uses spaces or underscores.
    """

    available_sheets = {
        normalized_name(sheet): sheet
        for sheet in sheet_names
    }

    for name in possible_names:

        normalized = normalized_name(name)

        if normalized in available_sheets:
            return available_sheets[normalized]

    return None


def normalize_ratio(value):
    """
    Convert ratios into one common format.

    Examples:
    1/1 -> 1:1
    2/1 -> 2:1
    1/1.5 -> 1:1.5
    """

    return (
        str(value)
        .strip()
        .upper()
        .replace(" ", "")
        .replace("/", ":")
    )


def display_ratio(value):
    """
    Display ratios using slash format.

    Example:
    2:1 -> 2/1
    """

    return normalize_ratio(value).replace(":", "/")


def normalize_gear_type(value):
    """
    Convert different gear-type descriptions into ST or SP.
    """

    value = str(value).strip().upper()

    mapping = {
        "STRAIGHT": "ST",
        "STRAIGHT BEVEL": "ST",
        "SPIRAL": "SP",
        "SPIRAL BEVEL": "SP",
    }

    return mapping.get(value, value)


def format_number(value):
    """
    Remove unnecessary zeros after a decimal.

    Example:
    90.0 -> 90
    75.50 -> 75.5
    """

    number = float(value)

    if number.is_integer():
        return f"{int(number):,}"

    return (
        f"{number:,.2f}"
        .rstrip("0")
        .rstrip(".")
    )


def ratio_sort_key(value):
    """
    Sort reduction ratios before increaser ratios.
    """

    ratio_text = normalize_ratio(value)

    ratio_match = re.match(
        r"^(\d+(?:\.\d+)?):(\d+(?:\.\d+)?)",
        ratio_text,
    )

    if not ratio_match:
        return 99, ratio_text

    numerator = float(ratio_match.group(1))
    denominator = float(ratio_match.group(2))

    direction_group = (
        0 if numerator >= denominator else 1
    )

    numerical_ratio = (
        numerator / denominator
        if denominator != 0
        else 999
    )

    return (
        direction_group,
        numerical_ratio,
        ratio_text,
    )


# =========================================================
# READ TORQUE DATABASE
# =========================================================

@st.cache_data
def load_database():
    """
    Read both worksheets:

    Torque_Range:
    Contains minimum and maximum torque ranges.

    Torque_Data:
    Contains detailed Input RPM and Output Torque ratings.
    """

    if not DATABASE_FILE.exists():

        raise FileNotFoundError(
            "HubCity_Bevel_Gear_Torque_Data.xlsx is missing. "
            "Keep the Excel database in the same folder "
            "as this Python file."
        )

    excel_file = pd.ExcelFile(
        DATABASE_FILE,
        engine="openpyxl",
    )

    # Find the torque-range worksheet
    range_sheet = find_sheet(
        excel_file.sheet_names,
        [
            "Torque_Range",
            "Torque_Ranges",
            "Torque Range",
            "Torque Ranges",
        ],
    )

    # Find the detailed RPM worksheet
    rpm_sheet = find_sheet(
        excel_file.sheet_names,
        [
            "Torque_Data",
            "Torque Data",
            "Raw_Data",
            "Raw Data",
        ],
    )

    if range_sheet is None:

        raise ValueError(
            "Torque_Range or Torque_Ranges worksheet "
            "was not found in the Excel database."
        )

    if rpm_sheet is None:

        raise ValueError(
            "Torque_Data worksheet was not found. "
            "The RPM rating table requires the detailed "
            "Torque_Data worksheet."
        )

    # Read both sheets
    range_source = pd.read_excel(
        DATABASE_FILE,
        sheet_name=range_sheet,
        engine="openpyxl",
    )

    rpm_source = pd.read_excel(
        DATABASE_FILE,
        sheet_name=rpm_sheet,
        engine="openpyxl",
    )

    # =====================================================
    # PREPARE TORQUE-RANGE DATA
    # =====================================================

    series_column = find_column(
        range_source,
        [
            "Series",
            "Model",
            "Gearbox Series",
        ],
    )

    ratio_column = find_column(
        range_source,
        [
            "Ratio",
            "Gear Ratio",
        ],
    )

    type_column = find_column(
        range_source,
        [
            "Gear_Type",
            "Gear Type",
            "Type",
        ],
    )

    minimum_column = find_column(
        range_source,
        [
            "Min_Torque_in_lb",
            "Minimum Torque",
            "Min Torque",
            "Minimum Torque in-lb",
        ],
    )

    maximum_column = find_column(
        range_source,
        [
            "Max_Torque_in_lb",
            "Maximum Torque",
            "Max Torque",
            "Maximum Torque in-lb",
        ],
    )

    direction_column = find_column(
        range_source,
        [
            "Direction",
        ],
    )

    required_range_columns = {
        "Series": series_column,
        "Ratio": ratio_column,
        "Gear Type": type_column,
        "Minimum Torque": minimum_column,
        "Maximum Torque": maximum_column,
    }

    missing_range_columns = [
        name
        for name, column
        in required_range_columns.items()
        if column is None
    ]

    if missing_range_columns:

        raise ValueError(
            "Missing torque-range columns: "
            + ", ".join(missing_range_columns)
        )

    torque_ranges = pd.DataFrame(
        {
            "Series": (
                range_source[series_column]
                .astype(str)
                .str.strip()
                .str.upper()
            ),
            "Ratio": (
                range_source[ratio_column]
                .apply(normalize_ratio)
            ),
            "Gear_Type": (
                range_source[type_column]
                .apply(normalize_gear_type)
            ),
            "Min_Torque": pd.to_numeric(
                range_source[minimum_column],
                errors="coerce",
            ),
            "Max_Torque": pd.to_numeric(
                range_source[maximum_column],
                errors="coerce",
            ),
        }
    )

    if direction_column is not None:

        torque_ranges["Direction"] = (
            range_source[direction_column]
            .astype(str)
            .str.strip()
        )

    else:

        torque_ranges["Direction"] = ""

    torque_ranges = torque_ranges[
        torque_ranges["Gear_Type"].isin(
            ["ST", "SP"]
        )
    ]

    torque_ranges = torque_ranges.dropna(
        subset=[
            "Min_Torque",
            "Max_Torque",
        ]
    )

    torque_ranges = torque_ranges.drop_duplicates(
        subset=[
            "Series",
            "Ratio",
            "Gear_Type",
        ],
        keep="first",
    )

    # =====================================================
    # PREPARE DETAILED RPM DATA
    # =====================================================

    rpm_series_column = find_column(
        rpm_source,
        [
            "Series",
            "Model",
            "Gearbox Series",
        ],
    )

    rpm_ratio_column = find_column(
        rpm_source,
        [
            "Ratio",
            "Gear Ratio",
        ],
    )

    rpm_type_column = find_column(
        rpm_source,
        [
            "Gear_Type",
            "Gear Type",
            "Type",
        ],
    )

    input_rpm_column = find_column(
        rpm_source,
        [
            "Input_RPM",
            "Input RPM",
            "RPM",
        ],
    )

    output_torque_column = find_column(
        rpm_source,
        [
            "Output_Torque_in_lb",
            "Output Torque",
            "Output Torque in-lb",
            "Torque",
        ],
    )

    torque_unit_column = find_column(
        rpm_source,
        [
            "Torque_Unit",
            "Torque Unit",
            "Unit",
        ],
    )

    catalogue_page_column = find_column(
        rpm_source,
        [
            "Catalogue_Page",
            "Catalogue Page",
            "Source Page",
            "Page",
        ],
    )

    required_rpm_columns = {
        "Series": rpm_series_column,
        "Ratio": rpm_ratio_column,
        "Gear Type": rpm_type_column,
        "Input RPM": input_rpm_column,
        "Output Torque": output_torque_column,
    }

    missing_rpm_columns = [
        name
        for name, column
        in required_rpm_columns.items()
        if column is None
    ]

    if missing_rpm_columns:

        raise ValueError(
            "Missing RPM-table columns: "
            + ", ".join(missing_rpm_columns)
        )

    rpm_ratings = pd.DataFrame(
        {
            "Series": (
                rpm_source[rpm_series_column]
                .astype(str)
                .str.strip()
                .str.upper()
            ),
            "Ratio": (
                rpm_source[rpm_ratio_column]
                .apply(normalize_ratio)
            ),
            "Gear_Type": (
                rpm_source[rpm_type_column]
                .apply(normalize_gear_type)
            ),
            "Input_RPM": pd.to_numeric(
                rpm_source[input_rpm_column],
                errors="coerce",
            ),
            "Output_Torque": pd.to_numeric(
                rpm_source[output_torque_column],
                errors="coerce",
            ),
        }
    )

    if torque_unit_column is not None:

        rpm_ratings["Torque_Unit"] = (
            rpm_source[torque_unit_column]
            .astype(str)
            .str.strip()
        )

    else:

        rpm_ratings["Torque_Unit"] = "in-lb"

    if catalogue_page_column is not None:

        rpm_ratings["Catalogue_Page"] = (
            rpm_source[catalogue_page_column]
            .astype(str)
            .str.strip()
        )

    else:

        rpm_ratings["Catalogue_Page"] = ""

    rpm_ratings = rpm_ratings[
        rpm_ratings["Gear_Type"].isin(
            ["ST", "SP"]
        )
    ]

    rpm_ratings = rpm_ratings.dropna(
        subset=[
            "Input_RPM",
            "Output_Torque",
        ]
    )

    rpm_ratings = rpm_ratings.drop_duplicates(
        subset=[
            "Series",
            "Ratio",
            "Gear_Type",
            "Input_RPM",
        ],
        keep="first",
    )

    return torque_ranges, rpm_ratings


# =========================================================
# CREATE EXCEL REPORT
# =========================================================

def create_excel_report(
    torque_ranges,
    selected_series,
):
    """
    Create one Excel workbook containing:

    1. Straight Bevel torque table
    2. Spiral Bevel torque table
    3. Selected Data worksheet
    """

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Torque Matrix"

    worksheet.sheet_view.showGridLines = False

    # Excel colours
    navy = "1F4E78"
    cyan = "5BC0DE"
    light_blue = "D9EAF7"
    unavailable_red = "F4CCCC"
    white = "FFFFFF"

    thin_side = Side(
        style="thin",
        color="404040",
    )

    cell_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    centre_alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    current_row = 1

    # Generate ST and SP tables
    for gear_type, table_title in [
        (
            "ST",
            "Straight Bevel Torque Range (in-lb)",
        ),
        (
            "SP",
            "Spiral Bevel Torque Range (in-lb)",
        ),
    ]:

        filtered_data = torque_ranges[
            torque_ranges["Series"].isin(
                selected_series
            )
            & (
                torque_ranges["Gear_Type"]
                == gear_type
            )
        ]

        available_ratios = sorted(
            filtered_data["Ratio"].unique(),
            key=ratio_sort_key,
        )

        torque_lookup = {
            (
                row["Series"],
                row["Ratio"],
            ): (
                f"{format_number(row['Min_Torque'])}"
                f" - "
                f"{format_number(row['Max_Torque'])}"
            )
            for _, row
            in filtered_data.iterrows()
        }

        last_column = max(
            2,
            len(available_ratios) + 1,
        )

        # Main table title
        worksheet.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=last_column,
        )

        title_cell = worksheet.cell(
            current_row,
            1,
            table_title,
        )

        title_cell.fill = PatternFill(
            "solid",
            fgColor=navy,
        )

        title_cell.font = Font(
            color=white,
            bold=True,
            size=13,
        )

        title_cell.alignment = centre_alignment

        # Header row
        header_row = current_row + 1

        worksheet.cell(
            header_row,
            1,
            "Series",
        )

        for column_number, ratio in enumerate(
            available_ratios,
            start=2,
        ):

            worksheet.cell(
                header_row,
                column_number,
                f"{display_ratio(ratio)} {gear_type}",
            )

        for cell in worksheet[
            header_row
        ][:last_column]:

            cell.fill = PatternFill(
                "solid",
                fgColor=cyan,
            )

            cell.font = Font(
                bold=True,
            )

            cell.alignment = centre_alignment
            cell.border = cell_border

        # Series data rows
        first_data_row = header_row + 1

        for row_number, series in enumerate(
            selected_series,
            start=first_data_row,
        ):

            series_cell = worksheet.cell(
                row_number,
                1,
                series,
            )

            series_cell.fill = PatternFill(
                "solid",
                fgColor=cyan,
            )

            series_cell.font = Font(
                bold=True,
            )

            series_cell.alignment = (
                centre_alignment
            )

            series_cell.border = cell_border

            for column_number, ratio in enumerate(
                available_ratios,
                start=2,
            ):

                torque_value = torque_lookup.get(
                    (
                        series,
                        ratio,
                    ),
                    "Not Available",
                )

                data_cell = worksheet.cell(
                    row_number,
                    column_number,
                    torque_value,
                )

                if torque_value == "Not Available":

                    data_cell.fill = PatternFill(
                        "solid",
                        fgColor=unavailable_red,
                    )

                else:

                    data_cell.fill = PatternFill(
                        "solid",
                        fgColor=light_blue,
                    )

                data_cell.alignment = (
                    centre_alignment
                )

                data_cell.border = cell_border

        current_row = (
            first_data_row
            + len(selected_series)
            + 2
        )

    # Column widths
    worksheet.column_dimensions["A"].width = 18

    for column_number in range(
        2,
        worksheet.max_column + 1,
    ):

        column_letter = get_column_letter(
            column_number
        )

        worksheet.column_dimensions[
            column_letter
        ].width = 18

    worksheet.freeze_panes = "B3"

    worksheet.page_setup.orientation = (
        "landscape"
    )

    worksheet.page_setup.fitToWidth = 1

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    # Detailed data worksheet
    details_sheet = workbook.create_sheet(
        "Selected Data"
    )

    details_sheet.append(
        [
            "Series",
            "Ratio",
            "Gear Type",
            "Direction",
            "Minimum Torque (in-lb)",
            "Maximum Torque (in-lb)",
        ]
    )

    selected_data = torque_ranges[
        torque_ranges["Series"].isin(
            selected_series
        )
    ].sort_values(
        [
            "Gear_Type",
            "Series",
            "Ratio",
        ]
    )

    for _, row in selected_data.iterrows():

        details_sheet.append(
            [
                row["Series"],
                display_ratio(row["Ratio"]),
                row["Gear_Type"],
                row["Direction"],
                row["Min_Torque"],
                row["Max_Torque"],
            ]
        )

    for cell in details_sheet[1]:

        cell.fill = PatternFill(
            "solid",
            fgColor=navy,
        )

        cell.font = Font(
            color=white,
            bold=True,
        )

        cell.alignment = centre_alignment

    details_sheet.freeze_panes = "A2"
    details_sheet.sheet_view.showGridLines = False

    details_widths = [
        18,
        15,
        15,
        15,
        24,
        24,
    ]

    for column_number, width in enumerate(
        details_widths,
        start=1,
    ):

        column_letter = get_column_letter(
            column_number
        )

        details_sheet.column_dimensions[
            column_letter
        ].width = width

    # Save Excel workbook into memory
    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output.getvalue()


# =========================================================
# LOAD DATABASE
# =========================================================

try:

    torque_ranges, rpm_ratings = (
        load_database()
    )

except Exception as error:

    st.error(str(error))

    st.stop()


# =========================================================
# CREATE APPLICATION TABS
# =========================================================

finder_tab, report_tab = st.tabs(
    [
        "🔎 Torque Finder",
        "📊 Excel Report Generator",
    ]
)


# =========================================================
# TAB 1: TORQUE FINDER
# =========================================================

with finder_tab:

    st.subheader(
        "Find Torque Range and RPM Ratings"
    )

    # Series selection
    selected_series = st.selectbox(
        "Series",
        sorted(
            torque_ranges[
                "Series"
            ].unique()
        ),
        key="finder_series",
    )

    series_data = torque_ranges[
        torque_ranges["Series"]
        == selected_series
    ]

    # Ratio selection
    selected_ratio = st.selectbox(
        "Ratio",
        sorted(
            series_data[
                "Ratio"
            ].unique(),
            key=ratio_sort_key,
        ),
        format_func=display_ratio,
        key="finder_ratio",
    )

    ratio_data = series_data[
        series_data["Ratio"]
        == selected_ratio
    ]

    # Gear type selection
    selected_type = st.selectbox(
        "Gear Type",
        sorted(
            ratio_data[
                "Gear_Type"
            ].unique()
        ),
        format_func=lambda value: (
            "ST - Straight Bevel"
            if value == "ST"
            else "SP - Spiral Bevel"
        ),
        key="finder_type",
    )

    final_result = ratio_data[
        ratio_data["Gear_Type"]
        == selected_type
    ]

    # Find Torque button
    if st.button(
        "Find Torque",
        type="primary",
    ):

        if final_result.empty:

            st.warning(
                "No matching catalogue data was found."
            )

        else:

            selected_record = (
                final_result.iloc[0]
            )

            minimum_torque = float(
                selected_record[
                    "Min_Torque"
                ]
            )

            maximum_torque = float(
                selected_record[
                    "Max_Torque"
                ]
            )

            # Display torque range
            st.success(
                "Available torque range: "
                f"{format_number(minimum_torque)} "
                "to "
                f"{format_number(maximum_torque)} "
                "in-lb"
            )

            minimum_column, maximum_column = (
                st.columns(2)
            )

            minimum_column.metric(
                "Minimum Torque",
                (
                    f"{format_number(minimum_torque)} "
                    "in-lb"
                ),
            )

            minimum_column.caption(
                (
                    f"{minimum_torque * IN_LB_TO_NM:,.2f} "
                    "N·m"
                )
            )

            maximum_column.metric(
                "Maximum Torque",
                (
                    f"{format_number(maximum_torque)} "
                    "in-lb"
                ),
            )

            maximum_column.caption(
                (
                    f"{maximum_torque * IN_LB_TO_NM:,.2f} "
                    "N·m"
                )
            )

            direction_value = str(
                selected_record["Direction"]
            ).strip()

            if (
                direction_value
                and direction_value.lower()
                != "nan"
            ):

                st.write(
                    f"Direction: **{direction_value}**"
                )

            # =============================================
            # RPM RATING TABLE
            # =============================================

            matching_rpm_data = rpm_ratings[
                (
                    rpm_ratings["Series"]
                    == selected_series
                )
                & (
                    rpm_ratings["Ratio"]
                    == selected_ratio
                )
                & (
                    rpm_ratings["Gear_Type"]
                    == selected_type
                )
            ].copy()

            st.subheader("RPM Rating Table")

            if matching_rpm_data.empty:

                st.warning(
                    "No detailed RPM rating rows were "
                    "found for this selection."
                )

            else:

                matching_rpm_data = (
                    matching_rpm_data.sort_values(
                        "Input_RPM",
                        ascending=False,
                    )
                )

                matching_rpm_data[
                    "Input RPM"
                ] = matching_rpm_data[
                    "Input_RPM"
                ].map(
                    lambda value:
                    format_number(value)
                )

                matching_rpm_data[
                    "Output Torque (in-lb)"
                ] = matching_rpm_data[
                    "Output_Torque"
                ].map(
                    lambda value:
                    format_number(value)
                )

                matching_rpm_data[
                    "Output Torque (N·m)"
                ] = (
                    matching_rpm_data[
                        "Output_Torque"
                    ]
                    * IN_LB_TO_NM
                ).round(2)

                rpm_display_table = (
                    matching_rpm_data[
                        [
                            "Input RPM",
                            "Output Torque (in-lb)",
                            "Output Torque (N·m)",
                            "Catalogue_Page",
                        ]
                    ].rename(
                        columns={
                            "Catalogue_Page":
                            "Catalogue Page",
                        }
                    )
                )

                st.dataframe(
                    rpm_display_table,
                    hide_index=True,
                    use_container_width=True,
                )


# =========================================================
# TAB 2: EXCEL REPORT GENERATOR
# =========================================================

with report_tab:

    st.subheader(
        "Generate ST and SP Excel Tables"
    )

    report_series = st.multiselect(
        "Which series do you want in the Excel report?",
        options=sorted(
            torque_ranges[
                "Series"
            ].unique()
        ),
        default=[],
    )

    if report_series:

        report_preview = torque_ranges[
            torque_ranges["Series"].isin(
                report_series
            )
        ]

        st.dataframe(
            report_preview[
                [
                    "Series",
                    "Ratio",
                    "Gear_Type",
                    "Min_Torque",
                    "Max_Torque",
                ]
            ],
            hide_index=True,
            use_container_width=True,
        )

        generated_excel = create_excel_report(
            torque_ranges,
            report_series,
        )

        st.download_button(
            "Generate and Download Excel",
            data=generated_excel,
            file_name=(
                "Hub_City_Bevel_Torque_"
                "ST_SP_Tables.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            type="primary",
            use_container_width=True,
        )

    else:

        st.info(
            "Select one or more gearbox series."
        )
        